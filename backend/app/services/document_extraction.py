from __future__ import annotations

from dataclasses import dataclass, replace
from difflib import SequenceMatcher
from io import BytesIO
from pathlib import Path
import math
import re
from typing import Callable, Protocol
import unicodedata
import zipfile
from xml.etree import ElementTree

from docx import Document
from PIL import Image, UnidentifiedImageError


MIN_TEXT_FOR_OCR = 30
HEADER_BAND = 0.06
FOOTER_BAND = 0.12
MAX_MARGINAL_LINE_CHARS = 180
DOCX_MAX_EMBEDDINGS = 20
DOCX_MAX_EMBEDDING_BYTES = 20 * 1024 * 1024
DOCX_MAX_EMBEDDINGS_TOTAL_BYTES = 50 * 1024 * 1024
DOCX_MAX_EMBEDDING_DEPTH = 2
PDF_MAX_EMBEDDINGS = 20
PDF_MAX_EMBEDDING_BYTES = 20 * 1024 * 1024
PDF_MAX_EMBEDDINGS_TOTAL_BYTES = 50 * 1024 * 1024
PAGE_NUMBER_PATTERN = re.compile(r"^(?:(?:page|p\.?\s*)\s*)?\d+(?:\s+(?:of|/)\s+\d+)?$", re.IGNORECASE)


@dataclass(frozen=True)
class _PDFTextLine:
    text: str
    top: float
    bottom: float


def _normalize_extracted_text(text: str) -> str:
    """Normalize layout noise while retaining line breaks and source wording."""
    normalized = unicodedata.normalize("NFKC", text).replace("\u00ad", "")
    normalized = normalized.replace("\r\n", "\n").replace("\r", "\n")
    cleaned_chars = []
    for character in normalized:
        if character == "\n":
            cleaned_chars.append(character)
        elif character == "\t" or character.isspace():
            cleaned_chars.append(" ")
        elif unicodedata.category(character) in {"Cc", "Cf"}:
            continue
        else:
            cleaned_chars.append(character)

    lines = [re.sub(r"\s+", " ", line).strip() for line in "".join(cleaned_chars).split("\n")]
    compact_lines: list[str] = []
    for line in lines:
        if line or (compact_lines and compact_lines[-1]):
            compact_lines.append(line)
    while compact_lines and not compact_lines[-1]:
        compact_lines.pop()
    return "\n".join(compact_lines)


def _marginal_signature(text: str) -> str:
    """Make page-number variants of the same running header compare equally."""
    normalized = _normalize_extracted_text(text).casefold()
    return re.sub(r"\d+", "#", normalized)


def _marginally_similar(left: str, right: str) -> bool:
    """Allow small OCR/format variations without matching short legal text."""
    left_signature = _marginal_signature(left)
    right_signature = _marginal_signature(right)
    if left_signature == right_signature:
        return True
    if min(len(left_signature), len(right_signature)) < 12:
        return False
    return SequenceMatcher(None, left_signature, right_signature, autojunk=False).ratio() >= 0.88


def _repeated_line_clusters(lines_by_page: list[list[str]]) -> list[str]:
    threshold = max(2, math.ceil(len(lines_by_page) * 0.6))
    clusters: list[tuple[str, set[int]]] = []
    for page_index, lines in enumerate(lines_by_page):
        for line in set(lines):
            if len(line) < 12 or len(line) > MAX_MARGINAL_LINE_CHARS:
                continue
            matching = next((cluster for cluster in clusters if _marginally_similar(cluster[0], line)), None)
            if matching is None:
                clusters.append((line, {page_index}))
            else:
                matching[1].add(page_index)
    return [signature for signature, indices in clusters if len(indices) >= threshold]


def _is_page_number(text: str) -> bool:
    return bool(PAGE_NUMBER_PATTERN.fullmatch(_normalize_extracted_text(text)))


def _pdf_text_lines(page) -> list[_PDFTextLine]:
    lines: list[_PDFTextLine] = []
    for block in page.get_text("dict").get("blocks", []):
        if block.get("type") != 0:
            continue
        for line in block.get("lines", []):
            text = _normalize_extracted_text("".join(span.get("text", "") for span in line.get("spans", [])))
            if text:
                bbox = line.get("bbox", (0, 0, 0, 0))
                lines.append(_PDFTextLine(text=text, top=float(bbox[1]), bottom=float(bbox[3])))
    return lines


def _repeated_pdf_marginalia(pages: list[list[_PDFTextLine]], page_heights: list[float]) -> list[list[str]]:
    """Find short headers/footers repeated in the geometric page margins."""
    if len(pages) < 2:
        return [[], []]

    top_by_page: list[list[str]] = [[] for _ in pages]
    bottom_by_page: list[list[str]] = [[] for _ in pages]
    for page_index, (lines, height) in enumerate(zip(pages, page_heights, strict=True)):
        for line in lines:
            if len(line.text) > MAX_MARGINAL_LINE_CHARS:
                continue
            if line.top <= height * HEADER_BAND:
                top_by_page[page_index].append(line.text)
            if line.bottom >= height * (1 - FOOTER_BAND):
                bottom_by_page[page_index].append(line.text)

    repeated_top = _repeated_line_clusters(top_by_page)
    repeated_bottom = _repeated_line_clusters(bottom_by_page)
    return [repeated_top, repeated_bottom]


def _clean_repeated_ocr_edges(units: list[ExtractedTextUnit]) -> list[ExtractedTextUnit]:
    """Remove repeated first/last OCR lines when they are clearly page furniture."""
    ocr_units = [unit for unit in units if "ocr" in unit.extraction_method.casefold()]
    if not ocr_units:
        return units

    unit_lines: dict[int, list[str]] = {}
    for page_index, unit in enumerate(ocr_units):
        lines = [line for line in _normalize_extracted_text(unit.text).splitlines() if line]
        unit_lines[page_index] = lines

    repeated_first = _repeated_line_clusters([[lines[0]] if len(lines) >= 2 else [] for lines in unit_lines.values()]) if len(ocr_units) >= 2 else []
    repeated_last = _repeated_line_clusters([[lines[-1]] if len(lines) >= 2 else [] for lines in unit_lines.values()]) if len(ocr_units) >= 2 else []
    cleaned: list[ExtractedTextUnit] = []
    ocr_index = 0
    for unit in units:
        if "ocr" not in unit.extraction_method.casefold():
            cleaned.append(unit)
            continue
        lines = unit_lines[ocr_index]
        ocr_index += 1
        if lines and (
            _is_page_number(lines[0])
            or any(_marginally_similar(lines[0], candidate) for candidate in repeated_first)
        ):
            lines = lines[1:]
        if lines and (
            _is_page_number(lines[-1])
            or any(_marginally_similar(lines[-1], candidate) for candidate in repeated_last)
        ):
            lines = lines[:-1]
        cleaned.append(replace(unit, text="\n".join(lines)))
    return cleaned


@dataclass(frozen=True)
class OCRResult:
    text: str
    confidence: float | None = None


@dataclass(frozen=True)
class ExtractedTextUnit:
    unit_number: int
    page_number: int | None
    locator_label: str
    text: str
    extraction_method: str
    ocr_confidence: float | None = None


@dataclass(frozen=True)
class ExtractionResult:
    units: list[ExtractedTextUnit]
    status: str
    error: str | None = None


class OCRBackend(Protocol):
    """Interface for local OCR and future optional cloud OCR providers."""

    name: str

    def recognize(self, image: Image.Image) -> OCRResult: ...


class TesseractOCR:
    name = "tesseract_ocr"

    def recognize(self, image: Image.Image) -> OCRResult:
        import pytesseract
        from pytesseract import Output
        from app.core.config import settings

        if settings.tesseract_cmd:
            pytesseract.pytesseract.tesseract_cmd = settings.tesseract_cmd

        data = pytesseract.image_to_data(
            image.convert("RGB"),
            lang="eng",
            output_type=Output.DICT,
            timeout=60,
        )
        lines: dict[tuple[int, int, int], list[str]] = {}
        for index, word in enumerate(data["text"]):
            if not word.strip():
                continue
            key = (data["block_num"][index], data["par_num"][index], data["line_num"][index])
            lines.setdefault(key, []).append(word.strip())
        confidences = [float(value) for value in data["conf"] if value not in {"-1", -1}]
        confidence = sum(confidences) / len(confidences) if confidences else None
        text = "\n".join(" ".join(words) for words in lines.values())
        return OCRResult(_normalize_extracted_text(text), confidence)


def _recognize_with_fallback(
    image: Image.Image,
    backends: list[OCRBackend],
) -> tuple[OCRResult | None, str | None, str | None]:
    errors: list[str] = []
    for backend in backends:
        try:
            result = backend.recognize(image)
            if result.text.strip():
                return result, backend.name, None
            errors.append(f"{backend.name} returned no text")
        except Exception as error:  # noqa: BLE001 — provider failures enable the fallback chain.
            errors.append(f"{backend.name}: {type(error).__name__}")
    if not backends:
        return None, None, "No OCR provider is configured."
    if any("TesseractNotFoundError" in error for error in errors):
        return (
            None,
            None,
            "Scanned content needs OCR. Install Tesseract OCR with its English language data, then retry.",
        )
    return None, None, "OCR could not read this scanned page. You can retry or review the original file."


def _ocr_quality_warnings(units: list[ExtractedTextUnit]) -> list[str]:
    from app.core.config import settings

    threshold = settings.ocr_review_confidence_threshold
    warnings = []
    for unit in units:
        if unit.ocr_confidence is not None and unit.ocr_confidence < threshold:
            warnings.append(f"{unit.locator_label} has low OCR confidence ({unit.ocr_confidence:.0f}%; review threshold {threshold:.0f}%). Verify against the original scan.")
        elif unit.ocr_confidence is None and "ocr" in unit.extraction_method.casefold():
            warnings.append(f"{unit.locator_label} has no OCR confidence score. Verify against the original scan.")
    return warnings


def _extract_pdf(source: Path | bytes, ocr_backends: list[OCRBackend], depth: int = 0) -> ExtractionResult:
    import fitz

    try:
        document = fitz.open(stream=source, filetype="pdf") if isinstance(source, bytes) else fitz.open(source)
    except Exception as error:
        raise ValueError("Could not read this PDF. It may be damaged or password-protected.") from error

    units: list[ExtractedTextUnit] = []
    errors: list[str] = []
    try:
        if document.needs_pass:
            raise ValueError("This PDF is password-protected and cannot be processed yet.")
        if document.page_count > 5000:
            raise ValueError("This PDF has more than the 5,000-page processing limit.")

        pdf_pages = list(document)
        pdf_lines = [_pdf_text_lines(page) for page in pdf_pages]
        marginalia = _repeated_pdf_marginalia(
            pdf_lines,
            [float(page.rect.height) for page in pdf_pages],
        )

        for page_index, (page, lines) in enumerate(zip(pdf_pages, pdf_lines, strict=True), start=1):
            top_headers, bottom_footers = marginalia
            retained_lines = []
            for line in lines:
                at_top = line.top <= page.rect.height * HEADER_BAND
                at_bottom = line.bottom >= page.rect.height * (1 - FOOTER_BAND)
                if (at_top and (any(_marginally_similar(line.text, header) for header in top_headers) or _is_page_number(line.text))) or (
                    at_bottom and (any(_marginally_similar(line.text, footer) for footer in bottom_footers) or _is_page_number(line.text))
                ):
                    continue
                retained_lines.append(line.text)
            text = _normalize_extracted_text("\n".join(retained_lines))
            method = "embedded_text"
            confidence = None
            if len(text) < MIN_TEXT_FOR_OCR:
                pixmap = page.get_pixmap(dpi=250, alpha=False)
                with Image.open(BytesIO(pixmap.tobytes("png"))) as image:
                    result, backend_name, ocr_error = _recognize_with_fallback(image, ocr_backends)
                if result is not None and len(result.text.strip()) >= len(text):
                    text = _normalize_extracted_text(result.text)
                    method = backend_name or "ocr"
                    confidence = result.confidence
                elif ocr_error:
                    errors.append(f"Page {page_index}: {ocr_error}")
            units.append(
                ExtractedTextUnit(
                    unit_number=page_index,
                    page_number=page_index,
                    locator_label=f"Page {page_index}",
                    text=text,
                    extraction_method=method,
                    ocr_confidence=confidence,
                )
            )

        # PDF portfolios, ordinary embedded files, and page-level file annotations
        # can all carry case material. Process both attachment forms under one set
        # of count/size/depth limits, and report anything that could not be read.
        try:
            attachment_count = document.embfile_count()
        except Exception:  # noqa: BLE001 — malformed PDFs may not expose an attachment catalog.
            attachment_count = 0
            errors.append("PDF embedded attachments could not be inspected; review the original PDF for attachments.")

        attachments: list[tuple[str, str, int, Callable[[], bytes]]] = []
        for attachment_index in range(attachment_count):
            try:
                info = document.embfile_info(attachment_index)
                filename = info.get("filename") or info.get("ufilename") or f"attachment-{attachment_index + 1}"
                filename = Path(str(filename)).name
                declared_size = int(info.get("size") or info.get("length") or 0)
                attachments.append((filename, f"Attachment {attachment_index + 1}", declared_size,
                                    lambda index=attachment_index: document.embfile_get(index)))
            except Exception as error:  # noqa: BLE001 — retain readable pages when one file spec is malformed.
                errors.append(f"Embedded PDF attachment {attachment_index + 1} could not be inspected ({type(error).__name__}).")

        annotation_count = 0
        for page_index, page in enumerate(pdf_pages, start=1):
            try:
                for annot in page.annots(types=[fitz.PDF_ANNOT_FILE_ATTACHMENT]) or []:
                    annotation_count += 1
                    info = annot.file_info
                    filename = info.get("filename") or info.get("ufilename") or f"page-{page_index}-attachment-{annotation_count}"
                    filename = Path(str(filename)).name
                    declared_size = int(info.get("size") or info.get("length") or 0)
                    attachments.append((filename, f"File attachment on page {page_index}", declared_size,
                                        lambda annotation=annot: annotation.get_file()))
            except Exception as error:  # noqa: BLE001 — annotation parsing is independent of page text extraction.
                errors.append(f"PDF file attachments on page {page_index} could not be inspected ({type(error).__name__}).")

        attachment_total_bytes = 0
        for attachment_index, (filename, source_label, declared_size, get_payload) in enumerate(
            attachments[:PDF_MAX_EMBEDDINGS], start=1
        ):
            try:
                if declared_size > PDF_MAX_EMBEDDING_BYTES:
                    errors.append(f"PDF {source_label.lower()} {filename} exceeds the 20 MB extraction limit and was skipped.")
                    continue
                payload = get_payload()
                attachment_total_bytes += len(payload)
                if len(payload) > PDF_MAX_EMBEDDING_BYTES or attachment_total_bytes > PDF_MAX_EMBEDDINGS_TOTAL_BYTES:
                    errors.append(f"PDF {source_label.lower()} {filename} exceeds the attachment size limit and was skipped.")
                    continue
                extension = _embedded_extension(payload, Path(filename).suffix.lower())
                if extension is None:
                    errors.append(f"PDF {source_label.lower()} {filename} uses an unsupported format and was not extracted.")
                    continue
                if depth >= DOCX_MAX_EMBEDDING_DEPTH:
                    errors.append(f"PDF {source_label.lower()} {filename} exceeds the nested attachment depth limit and was skipped.")
                    continue
                try:
                    nested_result = extract_document(payload, extension, ocr_backends, _depth=depth + 1)
                except ValueError as error:
                    errors.append(f"PDF {source_label.lower()} {filename} could not be extracted: {error}")
                    continue
                for nested_unit in nested_result.units:
                    if nested_unit.text.strip():
                        units.append(replace(
                            nested_unit,
                            unit_number=len(units) + 1,
                            locator_label=f"{source_label} · {filename} · {nested_unit.locator_label}"[:120],
                        ))
                if nested_result.error:
                    errors.append(f"PDF {source_label.lower()} {filename}: {nested_result.error}")
            except Exception as error:  # noqa: BLE001 — one bad attachment must not hide readable PDF pages.
                errors.append(f"PDF {source_label.lower()} {filename} could not be inspected ({type(error).__name__}).")
        if len(attachments) > PDF_MAX_EMBEDDINGS:
            errors.append(f"Only the first {PDF_MAX_EMBEDDINGS} of {len(attachments)} PDF attachments were checked.")
    finally:
        document.close()

    units = _clean_repeated_ocr_edges(units)
    errors.extend(_ocr_quality_warnings(units))
    error = " ".join(dict.fromkeys(errors)) or None
    status = "needs_review" if error or not any(unit.text for unit in units) else "processed"
    if not any(unit.text for unit in units) and error is None:
        error = "No readable text was found. The document may be blank or require OCR."
        status = "needs_review"
    return ExtractionResult(units, status, error)


def _extract_image(source: Path | bytes, ocr_backends: list[OCRBackend]) -> ExtractionResult:
    units: list[ExtractedTextUnit] = []
    errors: list[str] = []
    try:
        with Image.open(BytesIO(source) if isinstance(source, bytes) else source) as image:
            if getattr(image, "n_frames", 1) > 5000:
                raise ValueError("This image has more than the 5,000-frame processing limit.")
            for frame_index in range(image.n_frames):
                image.seek(frame_index)
                page = image.convert("RGB")
                result, backend_name, ocr_error = _recognize_with_fallback(page, ocr_backends)
                if result is None:
                    text = ""
                    confidence = None
                    method = "ocr_unavailable"
                    if ocr_error:
                        errors.append(f"Page {frame_index + 1}: {ocr_error}")
                else:
                    text = _normalize_extracted_text(result.text)
                    confidence = result.confidence
                    method = backend_name or "ocr"
                units.append(
                    ExtractedTextUnit(
                        unit_number=frame_index + 1,
                        page_number=frame_index + 1,
                        locator_label=f"Page {frame_index + 1}",
                        text=text,
                        extraction_method=method,
                        ocr_confidence=confidence,
                    )
                )
    except UnidentifiedImageError as error:
        raise ValueError("Could not read this image file.") from error

    units = _clean_repeated_ocr_edges(units)
    errors.extend(_ocr_quality_warnings(units))
    error = " ".join(dict.fromkeys(errors)) or None
    status = "needs_review" if error or not any(unit.text for unit in units) else "processed"
    return ExtractionResult(units, status, error)


def _extract_docx(source: Path | bytes, ocr_backends: list[OCRBackend], depth: int = 0) -> ExtractionResult:
    try:
        document = Document(BytesIO(source) if isinstance(source, bytes) else source)
    except Exception as error:
        raise ValueError("Could not read this DOCX file. It may be damaged or password-protected.") from error

    units: list[ExtractedTextUnit] = []
    for paragraph_number, paragraph in enumerate(document.paragraphs, start=1):
        text = _normalize_extracted_text(paragraph.text)
        if text:
            units.append(
                ExtractedTextUnit(
                    unit_number=len(units) + 1,
                    page_number=None,
                    locator_label=f"Paragraph {paragraph_number}",
                    text=text,
                    extraction_method="docx_text",
                )
            )

    for table_number, table in enumerate(document.tables, start=1):
        for row_number, row in enumerate(table.rows, start=1):
            for column_number, cell in enumerate(row.cells, start=1):
                text = _normalize_extracted_text("\n".join(paragraph.text for paragraph in cell.paragraphs))
                if text:
                    units.append(
                        ExtractedTextUnit(
                            unit_number=len(units) + 1,
                            page_number=None,
                            locator_label=f"Table {table_number}, row {row_number}, cell {column_number}",
                            text=text,
                            extraction_method="docx_text",
                        )
                    )

    seen_header_footer_text: set[tuple[str, str]] = set()
    for section_number, section in enumerate(document.sections, start=1):
        section_parts = (
            ("Header", section.header),
            ("First-page header", section.first_page_header),
            ("Even-page header", section.even_page_header),
            ("Footer", section.footer),
            ("First-page footer", section.first_page_footer),
            ("Even-page footer", section.even_page_footer),
        )
        for label, part in section_parts:
            part_paragraphs = [
                (f"{label}, section {section_number}, paragraph {paragraph_number}", paragraph.text)
                for paragraph_number, paragraph in enumerate(part.paragraphs, start=1)
            ]
            part_tables = [
                (f"{label}, section {section_number}, table {table_number}, row {row_number}, cell {column_number}",
                 "\n".join(paragraph.text for paragraph in cell.paragraphs))
                for table_number, table in enumerate(part.tables, start=1)
                for row_number, row in enumerate(table.rows, start=1)
                for column_number, cell in enumerate(row.cells, start=1)
            ]
            for locator, raw_text in part_paragraphs + part_tables:
                text = _normalize_extracted_text(raw_text)
                identity = ("header" if "header" in label.casefold() else "footer", text)
                if text and identity not in seen_header_footer_text:
                    seen_header_footer_text.add(identity)
                    units.append(ExtractedTextUnit(
                        unit_number=len(units) + 1,
                        page_number=None,
                        locator_label=locator,
                        text=text,
                        extraction_method="docx_header_footer",
                    ))

    warnings: list[str] = []
    package_parts = {
        "word/footnotes.xml": ("Footnote", "docx_footnote"),
        "word/endnotes.xml": ("Endnote", "docx_endnote"),
        "word/comments.xml": ("Comment", "docx_comment"),
    }
    try:
        with zipfile.ZipFile(BytesIO(source) if isinstance(source, bytes) else source) as package:
            for part_name, (label, method) in package_parts.items():
                if part_name not in package.namelist():
                    continue
                if package.getinfo(part_name).file_size > DOCX_MAX_EMBEDDING_BYTES:
                    warnings.append(f"{label} text exceeds the extraction size limit and was skipped.")
                    continue
                try:
                    root = ElementTree.fromstring(package.read(part_name))
                except ElementTree.ParseError:
                    warnings.append(f"{label} text could not be read from the DOCX package.")
                    continue
                tag = f"{{http://schemas.openxmlformats.org/wordprocessingml/2006/main}}"
                elements = list(root.findall(f"{tag}{label.lower()}"))
                for note in elements:
                    if label in {"Footnote", "Endnote"} and note.attrib.get(f"{tag}type") in {"separator", "continuationSeparator"}:
                        continue
                    text = _normalize_extracted_text("".join(node.text or "" for node in note.iter(f"{tag}t")))
                    if text:
                        identifier = note.attrib.get(f"{tag}id", str(len(units) + 1))
                        units.append(ExtractedTextUnit(
                            unit_number=len(units) + 1,
                            page_number=None,
                            locator_label=f"{label} {identifier}",
                            text=text,
                            extraction_method=method,
                        ))

            embedded = [name for name in package.namelist() if name.startswith("word/embeddings/") and not name.endswith("/")]
            if embedded:
                total_bytes = 0
                for attachment_index, member_name in enumerate(embedded[:DOCX_MAX_EMBEDDINGS], start=1):
                    info = package.getinfo(member_name)
                    total_bytes += info.file_size
                    if info.file_size > DOCX_MAX_EMBEDDING_BYTES or total_bytes > DOCX_MAX_EMBEDDINGS_TOTAL_BYTES:
                        warnings.append(f"Embedded attachment {Path(member_name).name} exceeds the extraction size limit and was skipped.")
                        continue
                    payload = package.read(member_name)
                    extension = _embedded_extension(payload, Path(member_name).suffix.lower())
                    if extension is None:
                        warnings.append(f"Embedded attachment {Path(member_name).name} uses an unsupported format and was not extracted.")
                        continue
                    if depth >= DOCX_MAX_EMBEDDING_DEPTH:
                        warnings.append(f"Embedded attachment {Path(member_name).name} exceeds the nested attachment depth limit and was skipped.")
                        continue
                    try:
                        nested_result = extract_document(payload, extension, ocr_backends, _depth=depth + 1)
                    except ValueError as error:
                        warnings.append(f"Embedded attachment {Path(member_name).name} could not be extracted: {error}")
                        continue
                    for nested_unit in nested_result.units:
                        if nested_unit.text.strip():
                            units.append(replace(
                                nested_unit,
                                unit_number=len(units) + 1,
                                locator_label=f"Embedded {attachment_index} · {Path(member_name).name} · {nested_unit.locator_label}"[:120],
                            ))
                    if nested_result.error:
                        warnings.append(f"Embedded attachment {Path(member_name).name}: {nested_result.error}")
                if len(embedded) > DOCX_MAX_EMBEDDINGS:
                    warnings.append(f"Only the first {DOCX_MAX_EMBEDDINGS} embedded attachments were checked.")
    except (OSError, zipfile.BadZipFile) as error:
        warnings.append("The DOCX package could not be checked for notes, comments, or embedded attachments.")

    if not units:
        return ExtractionResult([], "needs_review", "No readable paragraph, table, note, comment, or embedded attachment text was found." + (" " + " ".join(warnings) if warnings else ""))
    return ExtractionResult(units, "needs_review" if warnings else "processed", " ".join(warnings) or None)


def _embedded_extension(payload: bytes, declared_extension: str) -> str | None:
    if payload.startswith(b"%PDF-"):
        return ".pdf"
    if payload.startswith(b"\x89PNG\r\n\x1a\n"):
        return ".png"
    if payload.startswith(b"\xff\xd8\xff"):
        return ".jpg"
    if payload.startswith((b"II*\x00", b"MM\x00*", b"II+\x00", b"MM\x00+")):
        return ".tiff"
    if payload.startswith(b"PK\x03\x04"):
        try:
            with zipfile.ZipFile(BytesIO(payload)) as nested:
                if "word/document.xml" in nested.namelist():
                    return ".docx"
        except zipfile.BadZipFile:
            return None
    return declared_extension if declared_extension in {".pdf", ".docx", ".jpg", ".jpeg", ".png", ".tif", ".tiff"} else None


def extract_document(source: Path | bytes, extension: str, ocr_backends: list[OCRBackend] | None = None, *, _depth: int = 0) -> ExtractionResult:
    """Extract English text, preserving PDF/image pages and DOCX paragraph locations."""
    backends = ocr_backends if ocr_backends is not None else [TesseractOCR()]
    normalized_extension = extension.lower()
    if normalized_extension == ".pdf":
        return _extract_pdf(source, backends, _depth)
    if normalized_extension == ".docx":
        return _extract_docx(source, backends, _depth)
    if normalized_extension in {".jpg", ".jpeg", ".png", ".tif", ".tiff"}:
        return _extract_image(source, backends)
    raise ValueError("Unsupported file type for text extraction.")

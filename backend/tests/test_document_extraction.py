from __future__ import annotations

from io import BytesIO

import fitz
from docx import Document

from app.services.document_extraction import extract_document


def _text_pdf(text: str) -> bytes:
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), text)
    payload = document.tobytes()
    document.close()
    return payload


def _pdf_with_attachment(filename: str, payload: bytes) -> bytes:
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "The main pleading contains enough readable text for extraction.")
    document.embfile_add(filename, payload)
    result = document.tobytes()
    document.close()
    return result


def _pdf_with_file_annotation(filename: str, payload: bytes) -> bytes:
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "The filed application includes a separate exhibit for review.")
    page.add_file_annot((72, 100), payload, filename=filename)
    result = document.tobytes()
    document.close()
    return result


def test_pdf_embedded_pdf_attachment_is_extracted_with_source_locator(tmp_path) -> None:
    path = tmp_path / "filing.pdf"
    path.write_bytes(_pdf_with_attachment(
        "affidavit.pdf",
        _text_pdf("The witness affidavit confirms service on 12 March 2024."),
    ))

    result = extract_document(path, ".pdf", ocr_backends=[])

    attachment_units = [unit for unit in result.units if unit.locator_label.startswith("Attachment 1 · affidavit.pdf")]
    assert result.status == "processed"
    assert len(attachment_units) == 1
    assert "12 March 2024" in attachment_units[0].text


def test_pdf_unsupported_attachment_is_reported_as_needs_review(tmp_path) -> None:
    path = tmp_path / "filing.pdf"
    path.write_bytes(_pdf_with_attachment("exhibit.bin", b"unsupported attachment payload"))

    result = extract_document(path, ".pdf", ocr_backends=[])

    assert result.status == "needs_review"
    assert "exhibit.bin" in result.error
    assert "unsupported format" in result.error


def test_pdf_file_attachment_annotation_is_extracted_with_page_source_locator(tmp_path) -> None:
    path = tmp_path / "application.pdf"
    path.write_bytes(_pdf_with_file_annotation(
        "exhibit.pdf",
        _text_pdf("The attached exhibit records delivery on 21 April 2024."),
    ))

    result = extract_document(path, ".pdf", ocr_backends=[])

    attachment_units = [
        unit for unit in result.units
        if unit.locator_label.startswith("File attachment on page 1 · exhibit.pdf")
    ]
    assert result.status == "processed"
    assert len(attachment_units) == 1
    assert "21 April 2024" in attachment_units[0].text


def test_docx_header_and_footer_are_extracted_with_source_locations(tmp_path) -> None:
    document = Document()
    document.add_paragraph("The body of the pleading contains readable text.")
    document.sections[0].header.paragraphs[0].text = "Rao versus Mehta — Case No. 42/2024"
    document.sections[0].footer.paragraphs[0].text = "Confidential court filing"
    payload = BytesIO()
    document.save(payload)
    path = tmp_path / "pleading.docx"
    path.write_bytes(payload.getvalue())

    result = extract_document(path, ".docx", ocr_backends=[])

    header = next(unit for unit in result.units if unit.locator_label.startswith("Header"))
    footer = next(unit for unit in result.units if unit.locator_label.startswith("Footer"))
    assert "Case No. 42/2024" in header.text
    assert footer.text == "Confidential court filing"

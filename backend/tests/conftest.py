from collections.abc import Generator
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.core.config import settings
from app.core import auth
from app.core import durable_worker
from app.core.database import get_db
from app.db.base import Base
from app.main import app
from app.services import document_processing
from app.services import case_analysis_processing
from app.services import summary_processing


_TEST_EMAIL = "advocate@example.test"
_TEST_PASSWORD = "A-Strong-Test-Password-2026!"
_TEST_PASSWORD_HASH = auth.hash_password(_TEST_PASSWORD)


@pytest.fixture
def client(tmp_path, monkeypatch) -> Generator[TestClient, None, None]:
    monkeypatch.setattr(settings, "app_env", "test")
    # Keep synthetic uploads isolated from the developer's configured scanner.
    # Scanner behavior is covered separately with a deterministic test double.
    monkeypatch.setattr(settings, "clamav_command", None)
    with auth._failed_login_lock:
        auth._failed_login_times.clear()
    monkeypatch.setattr(settings, "advocate_email", _TEST_EMAIL)
    monkeypatch.setattr(settings, "advocate_password_hash", _TEST_PASSWORD_HASH)
    monkeypatch.setattr(settings, "session_secret", "test-only-session-secret-" * 2)
    monkeypatch.setattr(settings, "session_ttl_seconds", 3600)
    monkeypatch.setattr(settings, "login_attempt_limit", 5)
    monkeypatch.setattr(settings, "login_window_seconds", 900)
    test_engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    TestingSessionLocal = sessionmaker(bind=test_engine, autoflush=False, autocommit=False)
    monkeypatch.setattr(document_processing, "SessionLocal", TestingSessionLocal)
    monkeypatch.setattr(durable_worker, "SessionLocal", TestingSessionLocal)
    monkeypatch.setattr(case_analysis_processing, "SessionLocal", TestingSessionLocal)
    monkeypatch.setattr(summary_processing, "SessionLocal", TestingSessionLocal)
    Base.metadata.create_all(bind=test_engine)
    monkeypatch.setattr(settings, "uploads_dir", tmp_path / "uploads")

    def override_get_db():
        db = TestingSessionLocal()
        try:
            yield db
        finally:
            db.close()

    app.dependency_overrides[get_db] = override_get_db
    with TestClient(app) as test_client:
        test_client.headers.update({"Origin": settings.allowed_origins[0]})
        response = test_client.post(
            "/api/v1/auth/login",
            json={"email": _TEST_EMAIL, "password": _TEST_PASSWORD},
        )
        assert response.status_code == 200
        yield test_client
    app.dependency_overrides.clear()
    Base.metadata.drop_all(bind=test_engine)
    test_engine.dispose()


def make_docx() -> bytes:
    from docx import Document

    output = BytesIO()
    document = Document()
    document.add_paragraph("The agreement was signed on 12 March 2024.")
    document.add_paragraph("The witness confirmed delivery of the notice.")
    document.save(output)
    return output.getvalue()


@pytest.fixture
def docx_bytes() -> bytes:
    return make_docx()
